"""
Text extraction module for policy documents.

Supports PDF (PyMuPDF), DOCX (python-docx), and TXT formats.
Cleans extracted text while preserving legal numbering.
"""

import os
import re
from collections import Counter
from pathlib import Path


def extract_text(file_path: str) -> dict:
    """
    Extract and clean text from a policy document.

    Returns:
        {
            "text": str,           # full cleaned text
            "page_count": int,     # number of pages (1 for docx/txt)
            "word_count": int,
            "file_type": str,      # "pdf", "docx", "txt"
            "warnings": list[str]  # any issues encountered
        }
    """
    file_path = str(file_path)
    ext = Path(file_path).suffix.lower()

    if not os.path.isfile(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")

    extractors = {
        ".pdf": _extract_pdf,
        ".docx": _extract_docx,
        ".txt": _extract_txt,
    }

    extractor = extractors.get(ext)
    if extractor is None:
        raise ValueError(f"Unsupported file type: {ext}. Supported: pdf, docx, txt")

    raw_text, page_count, warnings = extractor(file_path)
    cleaned = _clean_text(raw_text)

    return {
        "text": cleaned,
        "page_count": page_count,
        "word_count": len(cleaned.split()),
        "file_type": ext.lstrip("."),
        "warnings": warnings,
    }


# ── Format-specific extractors ───────────────────────────────────────────


def _extract_pdf(file_path: str) -> tuple[str, int, list[str]]:
    """Extract text from PDF using PyMuPDF (fitz)."""
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz  # PyMuPDF < 1.25

    warnings = []
    pages_text = []

    try:
        doc = fitz.open(file_path)
    except Exception as exc:
        raise RuntimeError(f"Failed to open PDF: {exc}") from exc

    page_count = len(doc)
    empty_pages = 0

    for page_num, page in enumerate(doc, start=1):
        try:
            text = page.get_text("text")
            if not text.strip():
                empty_pages += 1
            pages_text.append(text)
        except Exception as exc:
            warnings.append(f"Page {page_num}: extraction failed ({exc})")
            pages_text.append("")

    doc.close()

    if empty_pages == page_count and page_count > 0:
        warnings.append(
            "All pages returned empty text. This PDF may be scanned/image-based. "
            "OCR is required for text extraction from scanned documents."
        )
    elif empty_pages > 0:
        warnings.append(
            f"{empty_pages} of {page_count} pages returned empty text "
            "(possibly scanned or image-only pages)."
        )

    return "\n\n".join(pages_text), page_count, warnings


def _extract_docx(file_path: str) -> tuple[str, int, list[str]]:
    """Extract text from DOCX including paragraphs and tables."""
    from docx import Document

    warnings = []

    try:
        doc = Document(file_path)
    except Exception as exc:
        raise RuntimeError(f"Failed to open DOCX: {exc}") from exc

    parts = []

    # Extract paragraphs
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    # Extract tables (policy clauses often live in tables)
    table_count = len(doc.tables)
    for table_idx, table in enumerate(doc.tables):
        table_rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                table_rows.append(" | ".join(cells))
        if table_rows:
            parts.append("\n".join(table_rows))

    if table_count > 0:
        warnings.append(f"Extracted text from {table_count} table(s).")

    return "\n\n".join(parts), 1, warnings


def _extract_txt(file_path: str) -> tuple[str, int, list[str]]:
    """Read plain text file with encoding fallback."""
    warnings = []

    for encoding in ("utf-8", "latin-1"):
        try:
            with open(file_path, "r", encoding=encoding) as f:
                text = f.read()
            if encoding != "utf-8":
                warnings.append(f"UTF-8 decode failed, used {encoding} fallback.")
            return text, 1, warnings
        except UnicodeDecodeError:
            continue

    raise RuntimeError(f"Cannot decode {file_path} with utf-8 or latin-1.")


# ── Text cleaning ────────────────────────────────────────────────────────


def _clean_text(text: str) -> str:
    """Clean extracted text while preserving legal numbering."""
    if not text:
        return ""

    # Remove repeated headers/footers
    text = _remove_repeated_headers(text)

    # Remove standalone page numbers (line with only digits, optionally with
    # surrounding whitespace or common decorators like "- 3 -" or "Page 3")
    text = re.sub(
        r"(?m)^\s*[-–—]?\s*(?:page\s*)?\d{1,4}\s*[-–—]?\s*$",
        "",
        text,
        flags=re.IGNORECASE,
    )

    # Strip non-meaningful special characters but PRESERVE:
    #   - Legal numbering: 1.2.3, (a), (ii), (iv), §5
    #   - Common punctuation: . , ; : ' " - / % & @ # !
    #   - Currency and legal symbols: ₹ $ € § ¶
    #   - Brackets: () [] {}
    text = re.sub(r"[^\w\s.,;:'\"\-/()\[\]{}&@#!%₹$€§¶°•·–—\n]", "", text)

    # Collapse excessive blank lines (3+ newlines become 2)
    text = re.sub(r"\n{3,}", "\n\n", text)

    # Collapse multiple spaces within a line
    text = re.sub(r"[^\S\n]{2,}", " ", text)

    # Strip leading/trailing whitespace per line
    lines = [line.strip() for line in text.splitlines()]
    text = "\n".join(lines)

    return text.strip()


def _remove_repeated_headers(text: str) -> str:
    """
    Detect and remove lines that repeat frequently across pages,
    likely headers or footers. A line repeating 3+ times and shorter
    than 120 chars is considered a header/footer.
    """
    lines = text.splitlines()

    if len(lines) < 10:
        return text

    # Count frequency of non-empty, short lines
    line_counts = Counter()
    for line in lines:
        stripped = line.strip()
        if stripped and len(stripped) < 120:
            line_counts[stripped] += 1

    # Lines appearing 3+ times are likely headers/footers
    repeated = {line for line, count in line_counts.items() if count >= 3}

    if not repeated:
        return text

    filtered = [
        line for line in lines
        if line.strip() not in repeated
    ]

    return "\n".join(filtered)


# ── CLI test ─────────────────────────────────────────────────────────────


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print("Usage: python text_extractor.py <file_path>")
        sys.exit(1)

    result = extract_text(sys.argv[1])
    print(f"Extracted {result['word_count']} words from {result['file_type']}")
    print(f"Pages: {result['page_count']}")
    print(f"Warnings: {result['warnings']}")
    print(f"\nFirst 500 chars:\n{result['text'][:500]}")
